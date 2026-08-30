#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Export GitHub organize Excel from live audit or existing audit JSON.

Usage:
  python export_report.py
  python export_report.py --from audit.json --out report.xlsx
  python export_report.py --keep a,b
"""
from __future__ import annotations

import argparse
import subprocess
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from common import (  # noqa: E402
    default_desktop,
    load_audit,
    run_audit,
    save_audit,
)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter


def row_fill(cat: str):
    fills = {
        "建议删除": PatternFill("solid", fgColor="FCE4EC"),
        "建议归档": PatternFill("solid", fgColor="FFF3E0"),
        "建议评估": PatternFill("solid", fgColor="FFF8E1"),
        "建议取消": PatternFill("solid", fgColor="FCE4EC"),
        "活跃": PatternFill("solid", fgColor="E8F5E9"),
        "保留": PatternFill("solid", fgColor="E3F2FD"),
        "已归档": PatternFill("solid", fgColor="ECEFF1"),
        "近期": PatternFill("solid", fgColor="E8F5E9"),
    }
    for k, v in fills.items():
        if k in (cat or ""):
            return v
    return None


def style_header(ws, cols):
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for i, c in enumerate(cols, 1):
        cell = ws.cell(1, i, c)
        cell.fill = fill
        cell.font = font


def autosize(ws, max_width=48):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col[:200]:
            val = "" if cell.value is None else str(cell.value)
            length = max(length, min(len(val), max_width))
        ws.column_dimensions[letter].width = max(12, min(max_width, length + 2))


def write_excel(out: Path, login: str, repo_rows: list, star_rows: list, keep: set[str]) -> None:
    wb = Workbook()
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    ws0 = wb.active
    ws0.title = "汇总说明"
    ws0["A1"] = f"GitHub 整理建议（{login}）"
    ws0["A1"].font = Font(bold=True, size=14)
    ws0["A2"] = f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    own_n = sum(1 for r in repo_rows if r["type"] == "自有")
    fork_n = sum(1 for r in repo_rows if r["type"] == "Fork")
    ws0["A3"] = f"仓库总数: {len(repo_rows)}（自有 {own_n} / Fork {fork_n}）"
    ws0["A4"] = f"星标总数: {len(star_rows)}"
    if keep:
        ws0["A5"] = "指定保留 Fork: " + ", ".join(sorted(keep))

    ws0["A7"] = "仓库处理分类统计"
    ws0["A7"].font = Font(bold=True)
    ws0["A8"] = "分类"
    ws0["B8"] = "数量"
    row_i = 9
    for k, v in sorted(Counter(r["cat"] for r in repo_rows).items(), key=lambda x: -x[1]):
        ws0[f"A{row_i}"] = k
        ws0[f"B{row_i}"] = v
        row_i += 1

    row_i += 1
    ws0[f"A{row_i}"] = "星标处理分类统计"
    ws0[f"A{row_i}"].font = Font(bold=True)
    row_i += 1
    ws0[f"A{row_i}"] = "分类"
    ws0[f"B{row_i}"] = "数量"
    row_i += 1
    for k, v in sorted(Counter(r["cat"] for r in star_rows).items(), key=lambda x: -x[1]):
        ws0[f"A{row_i}"] = k
        ws0[f"B{row_i}"] = v
        row_i += 1

    row_i += 1
    ws0[f"A{row_i}"] = "使用说明"
    ws0[f"A{row_i}"].font = Font(bold=True)
    for n in [
        "1. 「仓库」：自有仓 + Fork，含处理分类与建议；可用筛选按分类过滤。",
        "2. 「星标项目」：全部星标；建议取消/评估的排在前面；「建议归入List」用于建 GitHub Lists。",
        "3. 颜色：红粉=建议删除/取消；橙=建议归档；黄=需评估；绿=活跃/近期；蓝=保留；灰=已归档。",
        "4. 删除 Fork / 归档前请再确认；本表仅为建议。",
        "5. 检测由 scripts/audit.py / common.py 完成，勿手写重复检测逻辑。",
    ]:
        row_i += 1
        ws0[f"A{row_i}"] = n
    ws0.column_dimensions["A"].width = 80
    ws0.column_dimensions["B"].width = 12

    ws1 = wb.create_sheet("仓库")
    repo_headers = [
        "仓库名", "完整名", "类型", "可见性", "是否归档", "主语言",
        "Stars", "描述", "创建时间", "最后推送", "距今推送天数",
        "上游仓库", "领先提交ahead", "落后behind", "本人提交(默认分支)",
        "处理分类", "建议", "URL",
    ]
    style_header(ws1, repo_headers)
    for i, row in enumerate(repo_rows, 2):
        vals = [
            row["name"], row["full"], row["type"], row["vis"], row["archived"], row["lang"],
            row["stars"], row["desc"], row["created"], row["pushed"], row["days"],
            row["parent"], row["ahead"], row["behind"], row["user_c"],
            row["cat"], row["advice"], row["url"],
        ]
        for j, v in enumerate(vals, 1):
            cell = ws1.cell(i, j, v)
            cell.alignment = wrap
            cell.border = thin
        f = row_fill(row["cat"])
        if f:
            for j in range(1, len(vals) + 1):
                ws1.cell(i, j).fill = f
    ws1.auto_filter.ref = f"A1:R{len(repo_rows)+1}"
    ws1.freeze_panes = "A2"
    autosize(ws1)
    ws1.column_dimensions["H"].width = 40
    ws1.column_dimensions["Q"].width = 55

    ws2 = wb.create_sheet("星标项目")
    star_headers = [
        "仓库", "Owner", "主语言", "Topics", "描述",
        "上游Stars", "是否归档", "是否Fork",
        "最后推送", "距今推送天数", "星标时间", "距今星标天数",
        "处理分类", "建议归入List", "建议", "URL",
    ]
    style_header(ws2, star_headers)
    for i, row in enumerate(star_rows, 2):
        vals = [
            row["repo"], row["owner"], row["lang"], row["topics"], row["desc"],
            row["stars"], row["archived"], row["fork"],
            row["pushed"], row["pdays"], row["starred"], row["sdays"],
            row["cat"], row["list"], row["advice"], row["url"],
        ]
        for j, v in enumerate(vals, 1):
            cell = ws2.cell(i, j, v)
            cell.alignment = wrap
            cell.border = thin
        f = row_fill(row["cat"])
        if f:
            for j in range(1, len(vals) + 1):
                ws2.cell(i, j).fill = f
    ws2.auto_filter.ref = f"A1:P{len(star_rows)+1}"
    ws2.freeze_panes = "A2"
    autosize(ws2)
    ws2.column_dimensions["E"].width = 40
    ws2.column_dimensions["O"].width = 55

    wb.save(out)


def main() -> int:
    parser = argparse.ArgumentParser(description="Export GitHub organize report Excel")
    parser.add_argument("--from", dest="src", default="", help="Existing audit JSON from audit.py")
    parser.add_argument("--out", default="", help="Output .xlsx path")
    parser.add_argument("--audit-out", default="", help="Also save audit JSON when running live audit")
    parser.add_argument("--keep", default="", help="Comma-separated fork names to keep")
    parser.add_argument("--skip-fork-enrich", action="store_true")
    parser.add_argument("--skip-stars", action="store_true")
    args = parser.parse_args()
    keep = {x.strip() for x in args.keep.split(",") if x.strip()}

    if args.src:
        audit = load_audit(Path(args.src))
        if keep:
            # keep only affects classification; re-classify not done here — warn
            print("note: --keep ignored when --from is set (use keep at audit time)", flush=True)
        login = audit["summary"]["login"]
        repo_rows = audit["repos"]
        star_rows = audit["stars"]
        keep = set(audit["summary"].get("keep") or [])
    else:
        audit = run_audit(
            keep,
            skip_fork_enrich=args.skip_fork_enrich,
            skip_stars=args.skip_stars,
            progress=True,
        )
        login = audit["summary"]["login"]
        repo_rows = audit["repos"]
        star_rows = audit["stars"]
        if args.audit_out:
            save_audit(audit, Path(args.audit_out))
            print(f"audit json: {args.audit_out}", flush=True)

    out = Path(args.out) if args.out else (
        default_desktop() / f"GitHub整理建议_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    )
    out.parent.mkdir(parents=True, exist_ok=True)
    print(f"writing {out} ...", flush=True)
    write_excel(out, login, repo_rows, star_rows, keep)
    print(f"OK: {out}")
    print(f"repos={len(repo_rows)} stars={len(star_rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
