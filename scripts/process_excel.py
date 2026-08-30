#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""按 Excel「人工处理」列执行可自动化操作，并写回处理结果。

人工处理取值：
  - 空 / 不处理 → 跳过
  - 按照建议 → 按「处理分类」映射可自动动作并执行（删除类含星标上游）
  - 其它任意文字 → 归类意图；能自动则自动，否则写入进一步建议供 AI 分析
    常见：删除（不标上游）、删除并标星、归档、星标上游、取消星标

缺「人工处理」列的工作表会跳过（sheet_skipped），不中断其它表。

Usage:
  python process_excel.py --from report.xlsx --dry-run
  python process_excel.py --from report.xlsx --yes
  python process_excel.py --from report.xlsx --yes --out result.xlsx
"""
from __future__ import annotations

import argparse
import json
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))

from common import (  # noqa: E402
    ARCHIVE_OWN_CATEGORIES,
    DELETE_FORK_CATEGORIES,
    UNSTAR_CATEGORIES,
    archive_repo,
    auth_scopes,
    default_output_dir,
    delete_repo,
    ensure_starred,
    unstar_repo,
)

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

MANUAL_SKIP = "不处理"
MANUAL_FOLLOW = "按照建议"

# 结果列标题（导出与回写共用）
COL_MANUAL = "人工处理"
COL_RESULT = "处理结果"
COL_FURTHER = "进一步建议"

RESULT_SKIP = "跳过（不处理）"
RESULT_OK_PREFIX = "已执行"
RESULT_WOULD_PREFIX = "将执行"
RESULT_FAIL_PREFIX = "失败"
RESULT_NEED_AI = "需AI分析"
RESULT_NO_AUTO = "不可自动"


def _norm(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _header_map(ws) -> dict[str, int]:
    """列名 → 1-based 列号。"""
    m: dict[str, int] = {}
    for cell in ws[1]:
        if cell.value is not None:
            m[str(cell.value).strip()] = cell.column
    return m


def _ensure_result_cols(ws, headers: dict[str, int]) -> dict[str, int]:
    """若缺处理结果/进一步建议列则追加。"""
    max_col = max(headers.values()) if headers else 0
    for name in (COL_RESULT, COL_FURTHER):
        if name not in headers:
            max_col += 1
            cell = ws.cell(1, max_col, name)
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.font = Font(color="FFFFFF", bold=True)
            headers[name] = max_col
    if COL_MANUAL not in headers:
        raise ValueError(f"工作表「{ws.title}」缺少「{COL_MANUAL}」列，请用新版 export_report.py 导出")
    return headers


def _set_cell(ws, row: int, col: int, value: str) -> None:
    cell = ws.cell(row, col, value)
    cell.alignment = Alignment(wrap_text=True, vertical="top")


def classify_custom_intent(text: str, sheet: str) -> dict[str, Any]:
    """将「其他处理方式」文字归类为可识别意图。"""
    t = text.lower()
    cn = text

    # 明确跳过/保留
    if any(k in cn for k in ("不处理", "跳过", "忽略", "保留", "先不动", "暂不")):
        return {"intent": "skip", "auto": False, "note": "文字表示保留/跳过"}

    # 删除 fork / 仓库（「删除并标星」才星标上游；仅「删除」不标星）
    if any(k in cn for k in ("删除", "删掉", "删掉fork", "删 fork", "delete")):
        if sheet == "星标项目":
            return {
                "intent": "unstar",
                "auto": True,
                "note": "星标表中的「删除」按取消星标理解；若指删仓需到仓库表操作",
            }
        star_parent = any(k in cn for k in ("标星", "打星", "并星", "star"))
        return {
            "intent": "delete_fork",
            "auto": True,
            "note": "识别为删除（Fork）" + ("并星标上游" if star_parent else ""),
            "star_parent": star_parent,
        }

    # 归档
    if any(k in cn for k in ("归档", "archive")):
        return {"intent": "archive", "auto": True, "note": "识别为归档自有仓"}

    # 取消星标
    if any(k in cn for k in ("取消星标", "取消 star", "unstar", "去掉星标")):
        return {"intent": "unstar", "auto": True, "note": "识别为取消星标"}

    # 仅星标上游
    if any(k in cn for k in ("星标上游", "打星上游", "给上游打星", "star 上游", "star上游")):
        return {"intent": "star_parent", "auto": True, "note": "识别为仅星标上游"}

    # 密钥 / 工作流 — 不可自动
    if any(k in cn for k in ("密钥", "secret", "孤儿", "workflow", "工作流", "actions", "codeql")):
        return {
            "intent": "hygiene_manual",
            "auto": False,
            "note": "密钥/工作流类操作不自动执行，需人工在 GitHub Settings 处理",
        }

    # Lists
    if any(k in cn for k in ("list", "列表", "归类", "归入")):
        return {
            "intent": "star_list",
            "auto": False,
            "note": "GitHub Lists 需在网页创建/加入，脚本无法自动建 List",
        }

    # 按建议
    if MANUAL_FOLLOW in cn or "按建议" in cn or "采纳建议" in cn:
        return {"intent": "follow_advice", "auto": True, "note": "文字等价于按照建议"}

    return {
        "intent": "unknown",
        "auto": False,
        "note": f"未能可靠归类：「{text[:80]}」",
    }


def advice_action_for_repo(cat: str, typ: str) -> dict[str, Any] | None:
    """仓库表「按照建议」→ 动作。"""
    if typ == "Fork" and cat in DELETE_FORK_CATEGORIES:
        return {"intent": "delete_fork", "star_parent": True, "note": f"分类 {cat} → 星标上游并删除 Fork"}
    if typ == "自有" and cat in ARCHIVE_OWN_CATEGORIES:
        return {"intent": "archive", "note": f"分类 {cat} → 归档"}
    if "建议删除" in (cat or "") and typ == "Fork":
        return {"intent": "delete_fork", "star_parent": True, "note": f"分类 {cat} → 星标上游并删除 Fork"}
    if "建议归档" in (cat or "") and typ == "自有":
        return {"intent": "archive", "note": f"分类 {cat} → 归档"}
    return None


def advice_action_for_star(cat: str) -> dict[str, Any] | None:
    if cat in UNSTAR_CATEGORIES or "建议取消" in (cat or ""):
        return {"intent": "unstar", "note": f"分类 {cat} → 取消星标"}
    return None


def advice_action_for_hygiene(cat: str) -> dict[str, Any] | None:
    """工作流与密钥：一律不自动破坏性操作。"""
    if not cat or cat in ("正常", "已归档-跳过"):
        return {"intent": "skip", "auto": False, "note": "分类无需处理"}
    return {
        "intent": "hygiene_manual",
        "auto": False,
        "note": f"分类 {cat}：密钥删除/修 CI 请人工在 Settings / Actions 处理",
    }


def execute_intent(
    intent: str,
    *,
    full: str,
    parent: str = "",
    dry_run: bool,
    star_parent: bool = False,
) -> tuple[str, str]:
    """执行意图，返回 (处理结果, 进一步建议)。"""
    parts: list[str] = []
    further = ""

    if intent == "skip":
        return RESULT_SKIP, ""

    if intent == "star_parent":
        if dry_run:
            return f"{RESULT_WOULD_PREFIX}: 星标上游 {parent or '(无)'}", ""
        if not parent:
            return f"{RESULT_FAIL_PREFIX}: 无上游仓库", "请核对 Fork 上游字段"
        st = ensure_starred(parent)
        return f"{RESULT_OK_PREFIX}: 星标上游={st}", ""

    if intent == "delete_fork":
        if star_parent and parent:
            if dry_run:
                parts.append(f"星标上游 {parent}")
            else:
                st = ensure_starred(parent)
                parts.append(f"星标上游={st}")
        if dry_run:
            parts.append(f"删除 {full}")
            return f"{RESULT_WOULD_PREFIX}: " + "; ".join(parts), ""
        if not full:
            return f"{RESULT_FAIL_PREFIX}: 缺少完整名", ""
        dl = delete_repo(full)
        parts.append(f"删除={dl}")
        status = RESULT_OK_PREFIX if dl == "deleted" else RESULT_FAIL_PREFIX
        if dl == "failed:need_delete_repo_scope":
            further = "需要 gh auth refresh -h github.com -s delete_repo"
        return f"{status}: " + "; ".join(parts), further

    if intent == "archive":
        if dry_run:
            return f"{RESULT_WOULD_PREFIX}: 归档 {full}", ""
        if not full:
            return f"{RESULT_FAIL_PREFIX}: 缺少完整名", ""
        ar = archive_repo(full)
        status = RESULT_OK_PREFIX if ar == "archived" else RESULT_FAIL_PREFIX
        return f"{status}: 归档={ar}", ""

    if intent == "unstar":
        if dry_run:
            return f"{RESULT_WOULD_PREFIX}: 取消星标 {full}", ""
        if not full:
            return f"{RESULT_FAIL_PREFIX}: 缺少仓库名", ""
        us = unstar_repo(full)
        status = RESULT_OK_PREFIX if us == "unstarred" else RESULT_FAIL_PREFIX
        return f"{status}: 取消星标={us}", ""

    if intent == "hygiene_manual":
        return (
            RESULT_NO_AUTO,
            "请在 GitHub 仓库 Settings → Secrets / Actions 中人工处理；删除密钥前再确认是否被 reusable workflow 或动态名引用",
        )

    if intent == "star_list":
        return RESULT_NO_AUTO, "请在 GitHub Stars → Lists 中创建列表并加入该仓库"

    return RESULT_NEED_AI, "请由 Agent 根据人工处理文字与仓库上下文给出可执行步骤"


def process_repo_sheet(ws, dry_run: bool, summary: dict) -> None:
    headers = _ensure_result_cols(ws, _header_map(ws))
    for r in range(2, ws.max_row + 1):
        manual = _norm(ws.cell(r, headers[COL_MANUAL]).value)
        if not manual or manual == MANUAL_SKIP:
            _set_cell(ws, r, headers[COL_RESULT], RESULT_SKIP)
            _set_cell(ws, r, headers[COL_FURTHER], "")
            summary["skipped"] += 1
            continue

        typ = _norm(ws.cell(r, headers.get("类型", 0)).value) if "类型" in headers else ""
        cat = _norm(ws.cell(r, headers.get("处理分类", 0)).value) if "处理分类" in headers else ""
        full = _norm(ws.cell(r, headers.get("完整名", 0)).value) if "完整名" in headers else ""
        parent = _norm(ws.cell(r, headers.get("上游仓库", 0)).value) if "上游仓库" in headers else ""
        name = _norm(ws.cell(r, headers.get("仓库名", 0)).value) if "仓库名" in headers else ""

        if manual == MANUAL_FOLLOW:
            action = advice_action_for_repo(cat, typ)
            if not action:
                result = RESULT_NO_AUTO
                further = (
                    f"分类「{cat}」无对应自动动作。"
                    "可自动：Fork 建议删除类→删 fork；自有建议归档→归档。"
                    "请改写人工处理为具体操作，或由 AI 解读建议列。"
                )
                _set_cell(ws, r, headers[COL_RESULT], result)
                _set_cell(ws, r, headers[COL_FURTHER], further)
                summary["needs_ai"].append(
                    {"sheet": "仓库", "repo": full or name, "manual": manual, "cat": cat, "further": further}
                )
                summary["no_auto"] += 1
                continue
            intent = action["intent"]
            star_parent = bool(action.get("star_parent"))
            note = action.get("note", "")
        else:
            classified = classify_custom_intent(manual, "仓库")
            if classified["intent"] == "follow_advice":
                action = advice_action_for_repo(cat, typ)
                if not action:
                    further = classified["note"] + "；但当前分类无可自动动作，需 AI 解读"
                    _set_cell(ws, r, headers[COL_RESULT], RESULT_NEED_AI)
                    _set_cell(ws, r, headers[COL_FURTHER], further)
                    summary["needs_ai"].append(
                        {"sheet": "仓库", "repo": full or name, "manual": manual, "cat": cat, "further": further}
                    )
                    continue
                intent = action["intent"]
                star_parent = bool(action.get("star_parent"))
                note = action.get("note", "")
            elif not classified.get("auto"):
                further = classified["note"] + f"。原文：{manual}"
                result = RESULT_NEED_AI if classified["intent"] == "unknown" else RESULT_NO_AUTO
                _set_cell(ws, r, headers[COL_RESULT], result)
                _set_cell(ws, r, headers[COL_FURTHER], further)
                summary["needs_ai"].append(
                    {
                        "sheet": "仓库",
                        "repo": full or name,
                        "manual": manual,
                        "intent": classified["intent"],
                        "further": further,
                    }
                )
                if result == RESULT_NEED_AI:
                    summary["needs_ai_n"] += 1
                else:
                    summary["no_auto"] += 1
                continue
            else:
                intent = classified["intent"]
                star_parent = bool(classified.get("star_parent", intent == "delete_fork"))
                note = classified.get("note", "")

        result, further = execute_intent(
            intent, full=full, parent=parent, dry_run=dry_run, star_parent=star_parent
        )
        if note and not further:
            further = note
        elif note:
            further = f"{note}；{further}"
        _set_cell(ws, r, headers[COL_RESULT], result)
        _set_cell(ws, r, headers[COL_FURTHER], further)
        item = {"sheet": "仓库", "repo": full or name, "manual": manual, "result": result}
        if result.startswith(RESULT_OK_PREFIX) or result.startswith(RESULT_WOULD_PREFIX):
            summary["auto_done"].append(item)
            summary["auto_n"] += 1
        elif result.startswith(RESULT_FAIL_PREFIX):
            summary["failed"].append(item)
            summary["fail_n"] += 1
        else:
            summary["needs_ai"].append({**item, "further": further})
            summary["needs_ai_n"] += 1


def process_star_sheet(ws, dry_run: bool, summary: dict) -> None:
    headers = _ensure_result_cols(ws, _header_map(ws))
    for r in range(2, ws.max_row + 1):
        manual = _norm(ws.cell(r, headers[COL_MANUAL]).value)
        if not manual or manual == MANUAL_SKIP:
            _set_cell(ws, r, headers[COL_RESULT], RESULT_SKIP)
            _set_cell(ws, r, headers[COL_FURTHER], "")
            summary["skipped"] += 1
            continue

        cat = _norm(ws.cell(r, headers.get("处理分类", 0)).value) if "处理分类" in headers else ""
        full = _norm(ws.cell(r, headers.get("仓库", 0)).value) if "仓库" in headers else ""
        list_hint = (
            _norm(ws.cell(r, headers.get("建议归入List", 0)).value) if "建议归入List" in headers else ""
        )

        if manual == MANUAL_FOLLOW:
            action = advice_action_for_star(cat)
            if not action:
                further = (
                    f"分类「{cat}」无自动取消星标动作。"
                    + (f" 建议归入 List：{list_hint}。" if list_hint else "")
                    + "建 List 需人工；或改写人工处理列。"
                )
                _set_cell(ws, r, headers[COL_RESULT], RESULT_NO_AUTO)
                _set_cell(ws, r, headers[COL_FURTHER], further)
                summary["needs_ai"].append(
                    {"sheet": "星标项目", "repo": full, "manual": manual, "cat": cat, "further": further}
                )
                summary["no_auto"] += 1
                continue
            intent = action["intent"]
            note = action.get("note", "")
        else:
            classified = classify_custom_intent(manual, "星标项目")
            if classified["intent"] == "follow_advice":
                action = advice_action_for_star(cat)
                if not action:
                    further = "等价按照建议，但分类无可自动动作；" + classified["note"]
                    _set_cell(ws, r, headers[COL_RESULT], RESULT_NEED_AI)
                    _set_cell(ws, r, headers[COL_FURTHER], further)
                    summary["needs_ai"].append(
                        {"sheet": "星标项目", "repo": full, "manual": manual, "further": further}
                    )
                    summary["needs_ai_n"] += 1
                    continue
                intent = action["intent"]
                note = action.get("note", "")
            elif not classified.get("auto"):
                further = classified["note"]
                if list_hint and classified["intent"] == "star_list":
                    further += f"；建议 List={list_hint}"
                result = RESULT_NEED_AI if classified["intent"] == "unknown" else RESULT_NO_AUTO
                _set_cell(ws, r, headers[COL_RESULT], result)
                _set_cell(ws, r, headers[COL_FURTHER], further)
                summary["needs_ai"].append(
                    {
                        "sheet": "星标项目",
                        "repo": full,
                        "manual": manual,
                        "intent": classified["intent"],
                        "further": further,
                    }
                )
                if result == RESULT_NEED_AI:
                    summary["needs_ai_n"] += 1
                else:
                    summary["no_auto"] += 1
                continue
            else:
                intent = classified["intent"]
                note = classified.get("note", "")

        result, further = execute_intent(intent, full=full, dry_run=dry_run)
        if note:
            further = f"{note}；{further}" if further else note
        _set_cell(ws, r, headers[COL_RESULT], result)
        _set_cell(ws, r, headers[COL_FURTHER], further)
        item = {"sheet": "星标项目", "repo": full, "manual": manual, "result": result}
        if result.startswith(RESULT_OK_PREFIX) or result.startswith(RESULT_WOULD_PREFIX):
            summary["auto_done"].append(item)
            summary["auto_n"] += 1
        elif result.startswith(RESULT_FAIL_PREFIX):
            summary["failed"].append(item)
            summary["fail_n"] += 1
        else:
            summary["needs_ai"].append({**item, "further": further})
            summary["needs_ai_n"] += 1


def process_hygiene_sheet(ws, dry_run: bool, summary: dict) -> None:
    headers = _ensure_result_cols(ws, _header_map(ws))
    for r in range(2, ws.max_row + 1):
        manual = _norm(ws.cell(r, headers[COL_MANUAL]).value)
        if not manual or manual == MANUAL_SKIP:
            _set_cell(ws, r, headers[COL_RESULT], RESULT_SKIP)
            _set_cell(ws, r, headers[COL_FURTHER], "")
            summary["skipped"] += 1
            continue

        cat = _norm(ws.cell(r, headers.get("处理分类", 0)).value) if "处理分类" in headers else ""
        full = _norm(ws.cell(r, headers.get("仓库", 0)).value) if "仓库" in headers else ""
        orphans = (
            _norm(ws.cell(r, headers.get("孤儿Secrets列表", 0)).value)
            if "孤儿Secrets列表" in headers
            else ""
        )
        failed_wf = (
            _norm(ws.cell(r, headers.get("失败工作流名", 0)).value)
            if "失败工作流名" in headers
            else ""
        )
        advice = _norm(ws.cell(r, headers.get("建议", 0)).value) if "建议" in headers else ""

        if manual == MANUAL_FOLLOW:
            action = advice_action_for_hygiene(cat)
        else:
            classified = classify_custom_intent(manual, "工作流与密钥")
            if classified["intent"] == "follow_advice":
                action = advice_action_for_hygiene(cat)
            else:
                action = {
                    "intent": classified["intent"] if classified["intent"] != "unknown" else "hygiene_manual",
                    "auto": False,
                    "note": classified["note"],
                }

        # 密钥/CI 永不自动删除
        further_bits = [action.get("note", "")]
        if orphans:
            further_bits.append(f"孤儿密钥: {orphans}")
        if failed_wf:
            further_bits.append(f"失败工作流: {failed_wf}")
        if advice:
            further_bits.append(f"原建议: {advice}")
        if manual not in (MANUAL_FOLLOW, MANUAL_SKIP):
            further_bits.append(f"人工说明: {manual}")
        further_bits.append("请 Agent 根据上述信息给出逐步操作清单（不自动删 Secret）")
        further = "；".join(x for x in further_bits if x)

        result = RESULT_NO_AUTO if action.get("intent") != "unknown" else RESULT_NEED_AI
        if action.get("intent") == "skip":
            result = RESULT_SKIP
            further = action.get("note", "")
            summary["skipped"] += 1
        else:
            summary["needs_ai"].append(
                {
                    "sheet": "工作流与密钥",
                    "repo": full,
                    "manual": manual,
                    "cat": cat,
                    "further": further,
                    "orphans": orphans,
                    "failed_workflows": failed_wf,
                }
            )
            if result == RESULT_NEED_AI:
                summary["needs_ai_n"] += 1
            else:
                summary["no_auto"] += 1

        _set_cell(ws, r, headers[COL_RESULT], result)
        _set_cell(ws, r, headers[COL_FURTHER], further)
        _ = dry_run  # hygiene 不执行网络变更


def process_workbook(src: Path, out: Path, dry_run: bool) -> dict:
    wb = load_workbook(src)
    summary: dict[str, Any] = {
        "dry_run": dry_run,
        "source": str(src),
        "output": str(out),
        "skipped": 0,
        "auto_n": 0,
        "fail_n": 0,
        "no_auto": 0,
        "needs_ai_n": 0,
        "auto_done": [],
        "failed": [],
        "needs_ai": [],
        "processed_at": datetime.now().isoformat(timespec="seconds"),
    }

    if not dry_run:
        scopes = auth_scopes()
        # 仅当有删除意图时再强校验；此处先记录
        summary["auth_scopes"] = scopes

    def _has_manual(name: str) -> bool:
        if name not in wb.sheetnames:
            return False
        return COL_MANUAL in _header_map(wb[name])

    if _has_manual("仓库"):
        process_repo_sheet(wb["仓库"], dry_run, summary)
    elif "仓库" in wb.sheetnames:
        summary.setdefault("sheet_skipped", []).append("仓库:缺少人工处理列")

    if _has_manual("星标项目"):
        process_star_sheet(wb["星标项目"], dry_run, summary)
    elif "星标项目" in wb.sheetnames:
        summary.setdefault("sheet_skipped", []).append("星标项目:缺少人工处理列")

    if _has_manual("工作流与密钥"):
        process_hygiene_sheet(wb["工作流与密钥"], dry_run, summary)
    elif "工作流与密钥" in wb.sheetnames:
        summary.setdefault("sheet_skipped", []).append("工作流与密钥:缺少人工处理列")

    # 汇总说明追加处理摘要
    if "汇总说明" in wb.sheetnames:
        ws0 = wb["汇总说明"]
        row = (ws0.max_row or 1) + 2
        ws0[f"A{row}"] = "人工处理执行摘要"
        ws0[f"A{row}"].font = Font(bold=True)
        row += 1
        mode = "演练 dry-run" if dry_run else "已执行"
        for line in [
            f"模式: {mode}",
            f"时间: {summary['processed_at']}",
            f"跳过(不处理): {summary['skipped']}",
            f"自动成功/将执行: {summary['auto_n']}",
            f"失败: {summary['fail_n']}",
            f"不可自动: {summary['no_auto']}",
            f"需AI分析: {summary['needs_ai_n']}",
            "说明: 「需AI分析 / 不可自动」行见各表「进一步建议」列；Agent 应据此给出下一步。",
        ]:
            ws0[f"A{row}"] = line
            row += 1

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return summary


def main() -> int:
    p = argparse.ArgumentParser(description="Process GitHub organize Excel by 人工处理 column")
    p.add_argument("--from", dest="src", required=True, help="Filled Excel from export_report.py")
    p.add_argument("--out", default="", help="Output Excel with 处理结果 (default: beside source)")
    p.add_argument("--dry-run", action="store_true", help="Do not mutate GitHub")
    p.add_argument("--yes", action="store_true", help="Actually execute automatable actions")
    p.add_argument("--json-out", default="", help="Also write JSON summary path")
    args = p.parse_args()

    if not args.dry_run and not args.yes:
        p.error("Refusing to mutate without --yes (or pass --dry-run)")

    src = Path(args.src)
    if not src.is_file():
        print(json.dumps({"error": "file_not_found", "path": str(src)}, ensure_ascii=False))
        return 1

    if args.out:
        out = Path(args.out)
    else:
        stamp = datetime.now().strftime("%Y%m%d_%H%M")
        out = src.with_name(f"{src.stem}_处理结果_{stamp}.xlsx")
        if not out.parent.is_dir():
            out = default_output_dir() / out.name

    print(f"processing {src} -> {out} (dry_run={args.dry_run}) ...", flush=True)
    summary = process_workbook(src, out, dry_run=args.dry_run)

    # 精简打印：大列表保留
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    if args.json_out:
        Path(args.json_out).write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"json: {args.json_out}", flush=True)
    print(f"OK: {out}", flush=True)

    # 有失败时非零，便于 Agent 察觉
    if summary["fail_n"]:
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
